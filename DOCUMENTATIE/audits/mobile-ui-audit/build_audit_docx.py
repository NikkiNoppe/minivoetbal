#!/usr/bin/env python3
"""Bouw DOCUMENTATIE/audits/MOBILE_UI_AUDIT.docx uit de screenshots in deze map."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
DOCX_PATH = HERE.parent / "MOBILE_UI_AUDIT.docx"

NAVY = RGBColor(0x00, 0x5A, 0x94)

# (bestand, titel, route, rol, notitie)
SHOTS = [
    ("01-algemeen-admin.png", "Algemeen (start)", "/algemeen", "Publiek / ingelogd", "Eén h1 per pagina. Lightning-knop alleen bij admin-sessie."),
    ("05-competitie.png", "Competitie — stand", "/competitie", "Publiek", "Eerste speeldag: alle 0 punten (seizoen start)."),
    ("06-beker.png", "Beker — rounds", "/beker", "Publiek", "Horizontale round-nav + accordion."),
    ("07-playoff.png", "Play-off — nog niet gestart", "/playoff", "Publiek", "Lege-state card + CTA naar competitie."),
    ("08-reglement.png", "Reglement", "/reglement", "Publiek", "8 accordion-secties."),
    ("09-archief.png", "Archief", "/archief", "Publiek", "Seizoen-tab 2025-2026 + winnaars-cards."),
    ("35-nav-publiek-gast.png", "Mobiel menu — gast", "Sheet vanuit hamburger", "Gast", "Alleen Informatie + Inloggen."),
    ("36-modal-inloggen.png", "Modal — Inloggen", "Login-modal", "Gast", "Gebruikersnaam/e-mail + wachtwoord."),
    ("37-modal-wachtwoord-vergeten.png", "Modal — Wachtwoord vergeten", "Bovenop login", "Gast", "E-mail reset."),
    ("10-nav-admin-sheet.png", "Mobiel menu — admin (overzicht)", "Sheet", "Admin", "Profielkaart + groepen."),
    ("11-nav-informatie.png", "Menu — Informatie", "Sheet", "Admin", "Publieke pagina's vanuit admin-sessie."),
    ("12-nav-beheer.png", "Menu — Beheer (deels)", "Sheet", "Admin", "Beheer-groep onder de vouw; scroll nodig."),
    ("02-nav-admin-menu.png", "Menu — Wedstrijdformulieren", "Sheet", "Admin", "Competitie / Beker / Play-off."),
    ("03-nav-admin-beheer.png", "Menu — Beheer (volledig)", "Sheet", "Admin", "Spelers, Teams, Scheids, Users, Schorsingen, Financieel."),
    ("04-nav-admin-organisatie.png", "Menu — Organisatie", "Sheet", "Admin", "Blog, Berichten, Instellingen."),
    ("34-modal-snelle-admin-acties.png", "Modal — Snelle Acties (bliksem)", "Header-bliksem", "Admin", "Snelkoppelingen matchday / beheer / org."),
    ("13-match-forms-league.png", "Competitieformulieren", "/admin/match-forms/league", "Team / Admin", "Speeldag-accordions, Open-status."),
    ("14-modal-wedstrijdformulier-score.png", "Modal — Wedstrijdformulier (Score)", "Match klikken", "Team / Admin / Scheids", "Tabs Score / Spelers / Overig."),
    ("15-modal-wedstrijdformulier-spelers.png", "Modal — Wedstrijdformulier (Spelers)", "Tab Spelers", "Team / Admin / Scheids", "Thuis + uit; rugnummers."),
    ("16-modal-wedstrijdformulier-overig.png", "Modal — Wedstrijdformulier (Overig)", "Tab Overig", "Team / Admin / Scheids", "Kaarten / Financieel / Notities (dicht)."),
    ("17-match-forms-cup.png", "Bekerformulieren", "/admin/match-forms/cup", "Team / Admin", "Zelfde lijstpatroon als competitie."),
    ("18-match-forms-playoffs.png", "Play-off formulieren", "/admin/match-forms/playoffs", "Team / Admin", "Lege-state: play-offs nog niet gestart."),
    ("19-admin-players.png", "Spelerslijst", "/admin/players", "Team / Admin", "291 spelers; naam truncatie op mobiel."),
    ("20-modal-speler-bewerken.png", "Modal — Speler bewerken", "Potlood op speler", "Team / Admin", "Voornaam, achternaam, geboortedatum."),
    ("21-admin-teams.png", "Teams", "/admin/teams", "Admin", "Teamnaam truncatie; + Nieuw team."),
    ("22-modal-nieuw-team.png", "Modal — Nieuw team", "+ Nieuw team", "Admin", "Opslaan disabled tot naam ingevuld."),
    ("23-admin-users.png", "Gebruikersbeheer", "/admin/users", "Admin", "Rol-badge 'Teamverantwoo…' afgekapt."),
    ("24-modal-gebruiker-toevoegen.png", "Modal — Nieuwe gebruiker", "+ Gebruiker toevoegen", "Admin", "Rol + optioneel team."),
    ("25-admin-schorsingen.png", "Schorsingen Beheer", "/admin/schorsingen", "Team / Admin", "Regels bovenaan; 0 resultaten in lijst."),
    ("26-admin-scheidsrechters.png", "Scheidsrechter Beheer", "/admin/scheidsrechters", "Admin / Scheids", "Grid kan krap zijn op 390px."),
    ("27-admin-financial.png", "Financieel overzicht", "/admin/financial", "Admin", "Team-cards klikbaar."),
    ("28-modal-team-financial.png", "Modal — Team financieel detail", "Team-card klikken", "Admin", "Saldo + transacties; labels trunceren."),
    ("29-admin-settings.png", "Instellingen per competitie", "/admin/settings", "Admin", "Geen h1 — alleen h2 (a11y)."),
    ("30-admin-blog.png", "Blog beheer", "/admin/blog-management", "Admin", "Titel truncatie op live-post."),
    ("31-admin-notification.png", "Berichten", "/admin/notification", "Admin", "Recipient-lijst wrappen op mobiel."),
    ("32-profile.png", "Mijn Profiel", "/profile", "Ingelogd", "Accordions Berichten / Notities / Polls."),
    ("33-superadmin.png", "SuperAdmin platform", "/superadmin", "SuperAdmin (en dev)", "Tenant-keuze Harelbeke / Kuurne."),
    ("38-unsubscribe.png", "Uitschrijven (ongeldige link)", "/unsubscribe", "Publiek", "Foutstate zonder token."),
]


ROUTES = [
    ("Publiek", [
        ("/algemeen", "Start / nieuws", "Gast of ingelogd", "01"),
        ("/competitie", "Stand + schema", "Gast of ingelogd", "05"),
        ("/beker", "Bekerbracket", "Gast of ingelogd", "06"),
        ("/playoff", "Play-off stand", "Gast of ingelogd", "07"),
        ("/reglement", "Reglement", "Gast of ingelogd", "08"),
        ("/archief", "Historische seizoenen", "Gast of ingelogd", "09"),
        ("/kaarten", "Kaartenoverzicht", "Ingelogd", "Redirect → /algemeen (tab uit?)"),
        ("/reset-password", "Nieuw wachtwoord (token)", "Gast", "Zonder token → /algemeen"),
        ("/unsubscribe", "E-mail uitschrijven", "Gast", "38"),
    ]),
    ("Navigatie & auth", [
        ("Hamburger gast", "Publieke links + Inloggen", "Gast", "35"),
        ("Login-modal", "Sessie starten", "Gast", "36"),
        ("Wachtwoord vergeten", "Reset-mail", "Gast", "37"),
        ("Hamburger admin", "Volledige boom", "Admin", "10–12, 02–04"),
        ("Bliksem-menu", "Snelle acties", "Admin", "34"),
    ]),
    ("Ingelogd — formulieren", [
        ("/admin/match-forms/league", "Competitieformulieren", "Team / Admin", "13 + modal 14–16"),
        ("/admin/match-forms/cup", "Bekerformulieren", "Team / Admin", "17"),
        ("/admin/match-forms/playoffs", "Play-off formulieren", "Team / Admin", "18"),
        ("/profile", "Eigen account", "Elke rol", "32"),
    ]),
    ("Admin — beheer", [
        ("/admin/players", "Spelers", "Team / Admin", "19 + modal 20"),
        ("/admin/teams", "Teams", "Admin", "21 + modal 22"),
        ("/admin/users", "Gebruikers", "Admin", "23 + modal 24"),
        ("/admin/schorsingen", "Schorsingen (team/admin)", "Team / Admin", "25"),
        ("/admin/suspensions", "Schorsingen (admin-regels)", "Admin", "Zelfde UI als 25 of redirect"),
        ("/admin/scheidsrechters", "Scheids-rooster", "Admin / Scheids", "26"),
        ("/admin/financial", "Financieel", "Admin", "27 + modal 28"),
        ("/admin/settings", "Instellingen", "Admin", "29"),
        ("/admin/blog-management", "Blog", "Admin", "30"),
        ("/admin/notification", "Berichten", "Admin", "31"),
    ]),
    ("Admin — planning (redirect)", [
        ("/admin/competition", "Competitie beheer", "Admin", "Redirect → /admin/match-forms/league"),
        ("/admin/cup", "Beker beheer", "Admin", "Redirect → /admin/match-forms/league"),
        ("/admin/playoffs", "Play-off beheer", "Admin", "Redirect → /admin/match-forms/league"),
        ("/admin/season-calendar", "Seizoenskalender", "Admin", "Redirect → /admin/match-forms/league"),
        ("/admin/season-planning", "Seizoensplanning", "Admin", "Redirect → /admin/match-forms/league"),
        ("/admin/platform-beheer", "Platform beheer", "Admin", "Redirect → /admin/match-forms/league"),
    ]),
    ("SuperAdmin", [
        ("/superadmin", "Tenant-keuze", "SuperAdmin", "33"),
        ("/superadmin/:orgSlug", "Tenant-dashboard", "SuperAdmin", "Niet apart gefotografeerd"),
    ]),
]


def set_run_font(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def add_heading_styled(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, 9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shading = hdr[i]._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "005A94")
        shd.set(qn("w:val"), "clear")
        shading.append(shd)
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, 8)
    doc.add_paragraph()
    return table


def add_shot(doc, filename, title, route, role, note):
    path = HERE / filename
    add_heading_styled(doc, title, 2)
    meta = doc.add_paragraph()
    r = meta.add_run(f"Route: {route}   ·   Rol: {role}")
    set_run_font(r, 10, bold=True, color=NAVY)
    if note:
        n = doc.add_paragraph()
        rr = n.add_run(note)
        set_run_font(rr, 10)
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run()
        run.add_picture(str(path), width=Cm(7.2))
    else:
        miss = doc.add_paragraph()
        run = miss.add_run(f"[Screenshot ontbreekt: {filename}]")
        set_run_font(run, 10, color=RGBColor(0xB0, 0x00, 0x20))
    cap = doc.add_paragraph()
    cr = cap.add_run(filename)
    set_run_font(cr, 8, color=RGBColor(0x66, 0x66, 0x66))


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    title = doc.add_paragraph()
    tr = title.add_run("Mobile UI-audit")
    set_run_font(tr, 26, bold=True, color=NAVY)
    sub = doc.add_paragraph()
    sr = sub.add_run("Harelbeekse Minivoetbal  ·  7 september 2026  ·  v1.260907")
    set_run_font(sr, 12, color=NAVY)

    intro = doc.add_paragraph()
    ir = intro.add_run(
        "Overzicht van alle publieke en admin-pagina's plus de belangrijkste modals "
        "op mobiel (390×844, deviceScaleFactor 2). Gebruik dit document om gericht "
        "layout-, truncatie- en a11y-fouten op te ruimen. Screenshots zijn first-screen "
        "(boven de vouw), tenzij anders vermeld."
    )
    set_run_font(ir, 11)

    add_heading_styled(doc, "Hoe de routes volgen", 1)
    how = doc.add_paragraph()
    hr = how.add_run(
        "Start als gast op /algemeen → hamburger (Informatie) → login-modal. "
        "Na login als admin: hamburger-groepen Informatie → Wedstrijdformulieren → "
        "Beheer → Organisatie, of bliksem-menu. Match-kaart opent het wedstrijdformulier. "
        "Planning-URL's (/admin/competition, /cup, /playoffs, /season-calendar, "
        "/season-planning, /platform-beheer) redirecten nu naar /admin/match-forms/league "
        "(tab-visibility). /kaarten redirect naar /algemeen. SuperAdmin: /superadmin."
    )
    set_run_font(hr, 11)

    add_heading_styled(doc, "Eerste bevindingen (voor opruim)", 1)
    bullets = [
        "Admin-planningroutes redirecten naar competitieformulieren — geen screenshot van de echte beheer-UI.",
        "/kaarten (ingelogd) landt op /algemeen — tab waarschijnlijk uit of route-guard.",
        "/reset-password zonder token → /algemeen (verwacht); geen dedicated lege-token UI gefotografeerd.",
        "Instellingen: geen h1, alleen h2 «Instellingen per competitie».",
        "Truncatie op 390px: spelersnamen, teamnamen, gebruikersrollen, blogtitels, financiële labels.",
        "Scheids-grid is krap op mobiel (drie namen + statuscellen).",
        "Bliksem-knop in header op publieke pagina's wanneer admin ingelogd is — bewust?",
        "Dev-debugbalk is in deze set verborgen; productie heeft die balk niet.",
        "Gast-nav heeft geen Play-off/Beker-volgorde gelijk aan admin-Informatie (check consistentie).",
    ]
    for b in bullets:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(b)
        set_run_font(run, 11)

    add_heading_styled(doc, "Route-overzicht", 1)
    for zone, rows in ROUTES:
        add_heading_styled(doc, zone, 2)
        add_table(doc, ["Pad / actie", "Wat", "Rol", "Screenshot"], rows)

    add_heading_styled(doc, "Screenshots", 1)
    for item in SHOTS:
        add_shot(doc, *item)

    doc.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    main()
